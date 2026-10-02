"""gembot_core.doc_tools - Generating and reading Word (.docx) and Excel (.xlsx) documents."""
import os
from pathlib import Path


def create_docx(path: str, title: str, paragraphs: str) -> str:
    """Create a Microsoft Word (.docx) document with heading and paragraphs.

    Args:
        path: Output file path (e.g. assignment.docx)
        title: Document title or heading
        paragraphs: Multiline text separated by double newlines or lines
    """
    try:
        from docx import Document
        doc = Document()
        if title:
            doc.add_heading(title, 0)

        for p_text in paragraphs.split("\n\n"):
            p_text = p_text.strip()
            if p_text:
                if p_text.startswith("### "):
                    doc.add_heading(p_text[4:], level=3)
                elif p_text.startswith("## "):
                    doc.add_heading(p_text[3:], level=2)
                elif p_text.startswith("# "):
                    doc.add_heading(p_text[2:], level=1)
                elif p_text.startswith("- ") or p_text.startswith("* "):
                    for line in p_text.splitlines():
                        if line.startswith(("- ", "* ")):
                            doc.add_paragraph(line[2:], style='List Bullet')
                        else:
                            doc.add_paragraph(line)
                else:
                    doc.add_paragraph(p_text)

        p = Path(os.path.abspath(os.path.expanduser(os.path.expandvars(path))))
        p.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(p))
        return f"Successfully created Word document at {p}"
    except Exception as e:
        return f"Error creating docx: {e}"


def read_docx(path: str) -> str:
    """Read textual contents from a Word document (.docx)."""
    try:
        from docx import Document
        p = Path(os.path.abspath(os.path.expanduser(os.path.expandvars(path))))
        if not p.is_file():
            return f"Error: File '{p}' not found."
        doc = Document(str(p))
        text_lines = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(text_lines)
    except Exception as e:
        return f"Error reading docx: {e}"


def create_excel(path: str, sheet_name: str = "Sheet1", rows_csv: str = "") -> str:
    """Create a Microsoft Excel (.xlsx) spreadsheet from CSV-like data.

    Args:
        path: Destination file (e.g. data.xlsx)
        sheet_name: Title of the active worksheet
        rows_csv: Comma or pipe-separated lines representing rows and columns
    """
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        for row_str in rows_csv.strip().splitlines():
            row_str = row_str.strip()
            if not row_str:
                continue
            # Support both commas and pipes
            delimiter = "|" if "|" in row_str else ","
            cells = [c.strip() for c in row_str.split(delimiter)]
            ws.append(cells)

        p = Path(os.path.abspath(os.path.expanduser(os.path.expandvars(path))))
        p.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(p))
        return f"Successfully created Excel spreadsheet at {p}"
    except Exception as e:
        return f"Error creating Excel file: {e}"


def read_excel(path: str, max_rows: int = 50) -> str:
    """Read rows and columns from an Excel spreadsheet (.xlsx)."""
    try:
        import openpyxl
        p = Path(os.path.abspath(os.path.expanduser(os.path.expandvars(path))))
        if not p.is_file():
            return f"Error: File '{p}' not found."
        wb = openpyxl.load_workbook(str(p), data_only=True)
        ws = wb.active
        lines = []
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            if i > max_rows:
                lines.append(f"... [truncated after {max_rows} rows]")
                break
            filtered = [str(c) if c is not None else "" for c in row]
            if any(filtered):
                lines.append(" | ".join(filtered))
        return "\n".join(lines)
    except Exception as e:
        return f"Error reading Excel file: {e}"
